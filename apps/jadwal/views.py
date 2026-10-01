from datetime import date, datetime, time, timedelta
from io import BytesIO
from math import ceil

from django.contrib import messages
from django.core.exceptions import ValidationError
from django.db import transaction
from django.db.models import Q
from django.db.models.deletion import ProtectedError
from django.http import HttpResponse, JsonResponse
from django.shortcuts import get_object_or_404, redirect
from django.urls import reverse_lazy
from django.utils import timezone
from django.views.decorators.http import require_GET, require_POST
from django.views.generic import CreateView, DeleteView, DetailView, ListView, UpdateView

from apps.core.views import PostOnlyDeleteMixin
from apps.core.permissions import ADMIN_ROLE, ASISTEN_LAB_ROLE, LABORAN_ROLE
from apps.asleb.services import (
    get_active_asleb_for_pengguna,
    get_asleb_schedule_queryset,
)
from apps.kalender.realtime import send_schedule_change_request_update, send_schedule_update
from apps.pendaftaran_asleb.models import MataKuliahAsleb
from apps.ruangan.models import GrupRuanganGabungan, RuanganLab

from .forms import JadwalPraktikumForm
from .models import JadwalPraktikum, PermintaanPerubahanJadwal


@require_GET
def export_jadwal_praktikum_excel(request):
    schedules = list(
        JadwalPraktikum.objects.filter(status=JadwalPraktikum.STATUS_DITERIMA)
        .select_related('ruangan', 'ruangan_tambahan')
    )
    rooms = list(RuanganLab.objects.filter(aktif=True).order_by(
        '-kapasitas_tak_terbatas', 'nama'
    ))
    workbook = build_schedule_grid_xlsx(schedules, rooms)
    response = HttpResponse(
        workbook,
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    response['Content-Disposition'] = 'attachment; filename="jadwal-praktikum-labhub.xlsx"'
    return response


def build_schedule_grid_xlsx(schedules, rooms):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    slots = []
    cursor = datetime.combine(date.today(), time(7, 30))
    end = datetime.combine(date.today(), time(18, 0))
    while cursor < end:
        next_cursor = cursor + timedelta(minutes=30)
        slots.append((cursor.time(), next_cursor.time()))
        cursor = next_cursor

    workbook = Workbook()
    header_fill = PatternFill('solid', fgColor='CCFBF1')
    schedule_fill = PatternFill('solid', fgColor='E6FFFB')
    header_font = Font(name='Calibri', size=11, bold=True, color='0F766E')
    body_font = Font(name='Calibri', size=11, color='334155')
    schedule_font = Font(name='Calibri', size=11, bold=True, color='0F172A')
    thin_side = Side(style='thin', color='CBD5E1')
    grid_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

    for sheet_index, (day_key, day_label) in enumerate(JadwalPraktikum.HARI_CHOICES):
        worksheet = workbook.active if sheet_index == 0 else workbook.create_sheet()
        worksheet.title = day_label
        worksheet.freeze_panes = 'C2'
        worksheet.sheet_view.showGridLines = False
        worksheet.page_setup.orientation = 'landscape'
        worksheet.page_setup.fitToWidth = 1
        worksheet.page_setup.fitToHeight = 0

        day_schedules = [item for item in schedules if item.hari == day_key]
        headers = ['Dari', 'Sampai'] + [
            f'{room.nama} ({"Tak terbatas" if room.kapasitas_tak_terbatas else room.kapasitas})'
            for room in rooms
        ]
        worksheet.append(headers)
        worksheet.row_dimensions[1].height = 30
        for cell in worksheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.border = grid_border
            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

        for row_index, (slot_start, slot_end) in enumerate(slots, start=2):
            worksheet.cell(row=row_index, column=1, value=slot_start.strftime('%H:%M'))
            worksheet.cell(row=row_index, column=2, value=slot_end.strftime('%H:%M'))
            worksheet.row_dimensions[row_index].height = 32
            for column_index in range(1, len(headers) + 1):
                cell = worksheet.cell(row=row_index, column=column_index)
                cell.font = body_font
                cell.border = grid_border
                cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

        for item in day_schedules:
            start_index = next((i for i, pair in enumerate(slots, start=2) if pair[0] == item.waktu_mulai), None)
            if start_index is None:
                continue
            span = max(1, int((
                datetime.combine(date.today(), item.waktu_selesai)
                - datetime.combine(date.today(), item.waktu_mulai)
            ).total_seconds() // 1800))
            end_index = min(start_index + span - 1, len(slots) + 1)
            for room_id in {item.ruangan_id, item.ruangan_tambahan_id} - {None}:
                room_position = next((i for i, room in enumerate(rooms, start=3) if room.pk == room_id), None)
                if not room_position:
                    continue
                cell = worksheet.cell(row=start_index, column=room_position)
                cell.value = (
                    f'{item.mata_kuliah}\n{item.pengampu or "-"}\n{item.kelas or "-"}\n'
                    f'{item.waktu_mulai:%H:%M}-{item.waktu_selesai:%H:%M}'
                )
                cell.fill = schedule_fill
                cell.font = schedule_font
                cell.border = grid_border
                cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
                if end_index > start_index:
                    worksheet.merge_cells(
                        start_row=start_index,
                        start_column=room_position,
                        end_row=end_index,
                        end_column=room_position,
                    )
                    cell.border = grid_border

        worksheet.column_dimensions['A'].width = 12
        worksheet.column_dimensions['B'].width = 12
        for column_index in range(3, len(headers) + 1):
            worksheet.column_dimensions[get_column_letter(column_index)].width = 34
        worksheet.auto_filter.ref = f'A1:{get_column_letter(len(headers))}{len(slots) + 1}'

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def can_manage_jadwal(pengguna, jadwal):
    if not pengguna:
        return False
    if pengguna.role == LABORAN_ROLE:
        return True
    if pengguna.role == ASISTEN_LAB_ROLE:
        asleb = get_active_asleb_for_pengguna(pengguna)
        return bool(
            asleb
            and get_asleb_schedule_queryset(asleb).filter(pk=jadwal.pk).exists()
        )
    return False


class JadwalMutationAccessMixin:
    def dispatch(self, request, *args, **kwargs):
        pengguna = getattr(request, 'current_pengguna', None)
        if not pengguna or pengguna.role not in {LABORAN_ROLE, ASISTEN_LAB_ROLE}:
            messages.warning(request, 'Akses kelola jadwal hanya tersedia untuk Laboran dan Asisten Lab sesuai kewenangannya.')
            return redirect('jadwal:jadwal_list')

        return super().dispatch(request, *args, **kwargs)


class JadwalPraktikumListView(ListView):
    model = JadwalPraktikum
    template_name = 'jadwal/jadwal_list.html'
    context_object_name = 'jadwal_list'
    day_order = [key for key, _ in JadwalPraktikum.HARI_CHOICES]
    day_labels = dict(JadwalPraktikum.HARI_CHOICES)

    def get_selected_date(self):
        requested = self.request.GET.get('tanggal', '')
        try:
            selected = date.fromisoformat(requested) if requested else timezone.localdate()
        except ValueError:
            selected = timezone.localdate()
        requested_day = self.request.GET.get('hari', '').strip().lower()
        if not requested and requested_day in self.day_order:
            selected += timedelta(days=self.day_order.index(requested_day) - selected.weekday())
        return selected

    def get_selected_hari(self):
        if self.request.GET.get('tanggal'):
            weekday = self.get_selected_date().weekday()
            return self.day_order[weekday] if weekday < len(self.day_order) else 'senin'
        requested_hari = self.request.GET.get('hari', '').strip().lower()
        if requested_hari in self.day_order:
            return requested_hari

        today_index = timezone.localdate().weekday()
        if today_index < len(self.day_order):
            return self.day_order[today_index]

        return 'senin'

    def get_queryset(self):
        queryset = (
            JadwalPraktikum.objects.select_related('ruangan', 'ruangan_tambahan')
            .filter(
                hari=self.get_selected_hari(),
                status=JadwalPraktikum.STATUS_DITERIMA,
            )
            .order_by('waktu_mulai', 'ruangan__nama', 'mata_kuliah')
        )
        return queryset

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        selected_hari = self.get_selected_hari()
        selected_date = self.get_selected_date()
        if selected_date.weekday() == 6:
            selected_date += timedelta(days=1)
        ruangan_list = list(RuanganLab.objects.filter(aktif=True).order_by('nama'))
        context['current_pengguna'] = getattr(self.request, 'current_pengguna', None)
        context['hari_tabs'] = [
            {
                'value': value, 'label': label, 'active': value == selected_hari,
                'date': (selected_date - timedelta(days=selected_date.weekday())
                         + timedelta(days=index)).isoformat(),
            }
            for index, (value, label) in enumerate(JadwalPraktikum.HARI_CHOICES)
        ]
        context['selected_date'] = selected_date
        context['selected_hari'] = selected_hari
        context['selected_hari_label'] = self.day_labels[selected_hari]
        context['ruangan_list'] = ruangan_list
        time_slots, slot_keys = self.build_time_slots()
        context['time_slots'] = time_slots
        context['slot_count'] = len(time_slots)
        context['room_count'] = max(len(ruangan_list), 1)
        context['jadwal_blocks'] = self.build_jadwal_blocks(
            list(context['jadwal_list']),
            ruangan_list,
            slot_keys,
            context['current_pengguna'],
        )
        from apps.kalender.lab_events import conflicting_practicum_schedules
        from apps.kalender.views import get_visible_kegiatan_queryset

        context['lab_events'] = []
        for event in get_visible_kegiatan_queryset(context['current_pengguna']).select_related('ruangan').filter(
            tanggal=selected_date, ruangan__isnull=False,
        ):
            conflicts = list(conflicting_practicum_schedules(event))
            context['lab_events'].append({'event': event, 'conflicts': conflicts})
        conflicted_ids = {
            schedule.pk for item in context['lab_events'] for schedule in item['conflicts']
        }
        for block in context['jadwal_blocks']:
            block['event_conflict'] = block['jadwal'].pk in conflicted_ids
        context['lab_event_blocks'] = self.build_lab_event_blocks(
            [item['event'] for item in context['lab_events']],
            ruangan_list,
            slot_keys,
            selected_date,
        )
        context['praktikum_saya'] = self.get_praktikum_saya(context['current_pengguna'])
        if context['current_pengguna'] and context['current_pengguna'].role == LABORAN_ROLE:
            context['permintaan_perubahan'] = PermintaanPerubahanJadwal.objects.select_related(
                'jadwal', 'matkul', 'ruangan', 'ruangan_tambahan', 'diajukan_oleh'
            ).filter(status='diajukan')
        return context

    def get_praktikum_saya(self, pengguna):
        if not pengguna or pengguna.role != 'asisten_lab':
            return JadwalPraktikum.objects.none()

        asleb = get_active_asleb_for_pengguna(pengguna)
        if not asleb:
            return JadwalPraktikum.objects.none()

        return (
            get_asleb_schedule_queryset(
                asleb,
                statuses=[JadwalPraktikum.STATUS_DIAJUKAN, JadwalPraktikum.STATUS_DITERIMA],
            )
            .order_by('hari', 'waktu_mulai', 'ruangan__nama', 'mata_kuliah')
        )

    def build_time_slots(self):
        slots = []
        current_dt = datetime.combine(timezone.localdate(), time(7, 30))
        end_dt = datetime.combine(timezone.localdate(), time(18, 0))
        slot_keys = []

        while current_dt < end_dt:
            next_dt = current_dt + timedelta(minutes=30)
            mulai_label = current_dt.strftime('%H:%M')
            slot_keys.append(mulai_label)
            slots.append({
                'mulai': mulai_label.lstrip('0'),
                'selesai': next_dt.strftime('%H:%M').lstrip('0'),
            })
            current_dt = next_dt

        return slots, slot_keys

    def build_jadwal_blocks(self, jadwal_list, ruangan_list, slot_keys, pengguna):
        blocks = []
        ruangan_columns = {ruangan.pk: index + 1 for index, ruangan in enumerate(ruangan_list)}

        for jadwal in jadwal_list:
            start_key = self.get_slot_key(jadwal.waktu_mulai, slot_keys)
            occupied_columns = [
                ruangan_columns.get(room_id)
                for room_id in jadwal.get_occupied_room_ids()
                if ruangan_columns.get(room_id)
            ]
            grid_column = min(occupied_columns) if occupied_columns else None
            if not start_key or not grid_column:
                continue

            start_index = slot_keys.index(start_key)
            selesai = jadwal.waktu_selesai or (datetime.combine(timezone.localdate(), jadwal.waktu_mulai) + timedelta(minutes=30)).time()
            start_dt = datetime.combine(timezone.localdate(), jadwal.waktu_mulai)
            end_dt = datetime.combine(timezone.localdate(), selesai)
            if end_dt <= start_dt:
                end_dt = start_dt + timedelta(minutes=30)

            span = max(1, int((end_dt - start_dt).total_seconds() // 1800))
            span = min(span, len(slot_keys) - start_index)
            column_span = max(1, len(occupied_columns))
            blocks.append({
                'jadwal': jadwal,
                'grid_column': grid_column,
                'column_span': column_span,
                'grid_row': start_index + 1,
                'span': span,
                'can_manage': can_manage_jadwal(pengguna, jadwal),
            })

        return blocks

    def build_lab_event_blocks(self, events, ruangan_list, slot_keys, selected_date):
        if not slot_keys:
            return []
        room_columns = {room.pk: index + 1 for index, room in enumerate(ruangan_list)}
        board_start = datetime.combine(selected_date, time.fromisoformat(slot_keys[0]))
        board_end = board_start + timedelta(minutes=30 * len(slot_keys))
        blocks = []
        for event in events:
            column = room_columns.get(event.ruangan_id)
            if not column or not event.waktu_selesai:
                continue
            start = max(datetime.combine(selected_date, event.waktu_mulai), board_start)
            end = min(datetime.combine(selected_date, event.waktu_selesai), board_end)
            if end <= start:
                continue
            start_index = int((start - board_start).total_seconds() // 1800)
            end_index = ceil((end - board_start).total_seconds() / 1800)
            blocks.append({
                'event': event,
                'grid_column': column,
                'grid_row': start_index + 1,
                'span': end_index - start_index,
            })
        return blocks

    def get_slot_key(self, value, slot_keys):
        value_dt = datetime.combine(timezone.localdate(), value)
        selected_key = None

        for slot_key in slot_keys:
            slot_dt = datetime.combine(timezone.localdate(), datetime.strptime(slot_key, '%H:%M').time())
            if slot_dt <= value_dt:
                selected_key = slot_key
            else:
                break

        return selected_key


class JadwalPraktikumDetailView(DetailView):
    model = JadwalPraktikum
    template_name = 'jadwal/jadwal_detail.html'
    context_object_name = 'jadwal'

    def get_queryset(self):
        queryset = super().get_queryset()
        pengguna = getattr(self.request, 'current_pengguna', None)
        if pengguna and pengguna.role == 'mahasiswa':
            return queryset.filter(status=JadwalPraktikum.STATUS_DITERIMA)
        if pengguna and pengguna.role == ADMIN_ROLE:
            return queryset.filter(status=JadwalPraktikum.STATUS_DITERIMA)
        return queryset

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['current_pengguna'] = getattr(self.request, 'current_pengguna', None)
        context['can_manage_jadwal'] = can_manage_jadwal(context['current_pengguna'], self.object)
        context['permintaan_terakhir'] = self.object.permintaan_perubahan.select_related(
            'ruangan', 'ruangan_tambahan', 'diajukan_oleh', 'diproses_oleh'
        ).first()
        return context


class JadwalPraktikumCreateView(JadwalMutationAccessMixin, CreateView):
    model = JadwalPraktikum
    form_class = JadwalPraktikumForm
    template_name = 'jadwal/jadwal_form.html'
    success_url = reverse_lazy('jadwal:jadwal_list')

    def dispatch(self, request, *args, **kwargs):
        pengguna = getattr(request, 'current_pengguna', None)
        if pengguna and pengguna.role == LABORAN_ROLE:
            messages.warning(request, 'Laboran tidak dapat menambah jadwal praktikum. Gunakan Kalender untuk booking ruangan.')
            return redirect('jadwal:jadwal_list')
        return super().dispatch(request, *args, **kwargs)

    def get_form_kwargs(self):
        kwargs = super().get_form_kwargs()
        kwargs['current_pengguna'] = getattr(self.request, 'current_pengguna', None)
        return kwargs

    def form_valid(self, form):
        pengguna = getattr(self.request, 'current_pengguna', None)
        if pengguna and pengguna.role == 'asisten_lab':
            mata_kuliah = str(form.cleaned_data['matkul'])
            with transaction.atomic():
                duplicate_exists = JadwalPraktikum.objects.select_for_update().filter(
                    mata_kuliah=mata_kuliah,
                    status__in=[
                        JadwalPraktikum.STATUS_DIAJUKAN,
                        JadwalPraktikum.STATUS_DITERIMA,
                    ],
                ).exists()
                if duplicate_exists:
                    form.add_error('matkul', 'Jadwal untuk mata kuliah ini sudah pernah diajukan.')
                    return self.form_invalid(form)
                response = super().form_valid(form)
                transaction.on_commit(lambda: send_schedule_update(
                    self.object,
                    event='schedule.submitted',
                    notify_managers=True,
                ))
                return response

        response = super().form_valid(form)
        transaction.on_commit(lambda: send_schedule_update(self.object, event='schedule.created'))
        return response


class JadwalPraktikumUpdateView(JadwalMutationAccessMixin, UpdateView):
    model = JadwalPraktikum
    form_class = JadwalPraktikumForm
    template_name = 'jadwal/jadwal_form.html'
    success_url = reverse_lazy('jadwal:jadwal_list')

    def get_queryset(self):
        queryset = super().get_queryset()
        pengguna = getattr(self.request, 'current_pengguna', None)
        if pengguna and pengguna.role == ASISTEN_LAB_ROLE:
            asleb = get_active_asleb_for_pengguna(pengguna)
            if not asleb:
                return queryset.none()
            return queryset.filter(pk__in=get_asleb_schedule_queryset(asleb).values('pk'))
        if pengguna and pengguna.role == LABORAN_ROLE:
            return queryset
        return queryset.none()
        return queryset

    def get_form_kwargs(self):
        kwargs = super().get_form_kwargs()
        kwargs['current_pengguna'] = getattr(self.request, 'current_pengguna', None)
        return kwargs

    def form_valid(self, form):
        pengguna = getattr(self.request, 'current_pengguna', None)
        if pengguna and pengguna.role == 'asisten_lab':
            if self.object.permintaan_perubahan.filter(status='diajukan').exists():
                messages.warning(self.request, 'Masih ada permintaan perubahan jadwal yang menunggu persetujuan laboran.')
                return redirect('jadwal:jadwal_detail', pk=self.object.pk)
            change_request = PermintaanPerubahanJadwal.objects.create(
                jadwal=self.object,
                matkul=form.cleaned_data['matkul'],
                ruangan=form.cleaned_data['ruangan'],
                ruangan_tambahan=form.cleaned_data.get('ruangan_tambahan'),
                hari=form.cleaned_data['hari'],
                waktu_mulai=form.cleaned_data['waktu_mulai'],
                waktu_selesai=form.cleaned_data.get('waktu_selesai'),
                catatan=form.cleaned_data.get('catatan', ''),
                diajukan_oleh=pengguna,
            )
            transaction.on_commit(lambda request_id=change_request.pk: send_schedule_change_request_update(
                PermintaanPerubahanJadwal.objects.select_related(
                    'jadwal', 'matkul', 'ruangan', 'ruangan_tambahan', 'diajukan_oleh'
                ).get(pk=request_id)
            ))
            messages.success(self.request, 'Permintaan perubahan jadwal dikirim dan menunggu persetujuan laboran.')
            return redirect('jadwal:jadwal_detail', pk=self.object.pk)
        response = super().form_valid(form)
        transaction.on_commit(lambda: send_schedule_update(self.object, event='schedule.updated'))
        return response


class JadwalPraktikumDeleteView(JadwalMutationAccessMixin, PostOnlyDeleteMixin, DeleteView):
    model = JadwalPraktikum
    template_name = 'jadwal/jadwal_confirm_delete.html'
    context_object_name = 'jadwal'
    success_url = reverse_lazy('jadwal:jadwal_list')

    def get_queryset(self):
        queryset = super().get_queryset()
        pengguna = getattr(self.request, 'current_pengguna', None)
        if pengguna and pengguna.role == ASISTEN_LAB_ROLE:
            return queryset.none()
        if pengguna and pengguna.role == LABORAN_ROLE:
            return queryset
        return queryset

    def post(self, request, *args, **kwargs):
        self.object = self.get_object()
        try:
            response = super().post(request, *args, **kwargs)
            messages.success(
                request,
                'Jadwal praktikum berhasil dihapus. Riwayat absensi lama tetap disimpan tanpa referensi jadwal.'
            )
            return response
        except ProtectedError:
            related_absensi_count = self.object.absensi_asleb.count()
            related_absensi_masuk_count = self.object.absensi_masuk_asleb.count()
            total_references = related_absensi_count + related_absensi_masuk_count
            messages.error(
                request,
                'Jadwal praktikum tidak bisa dihapus karena masih dipakai oleh '
                f'{total_references} data absensi. Hapus atau pindahkan referensi absensinya terlebih dahulu.'
            )
            return redirect('jadwal:jadwal_detail', pk=self.object.pk)


@require_POST
def process_schedule_change_request(request, pk, decision):
    pengguna = getattr(request, 'current_pengguna', None)
    if not pengguna or pengguna.role != LABORAN_ROLE:
        messages.error(request, 'Hanya laboran yang dapat memproses perubahan jadwal.')
        return redirect('jadwal:jadwal_list')

    change_request = get_object_or_404(
        PermintaanPerubahanJadwal.objects.select_related('jadwal', 'matkul', 'ruangan', 'ruangan_tambahan'),
        pk=pk,
        status='diajukan',
    )
    if decision == 'tolak':
        change_request.status = 'ditolak'
        change_request.diproses_oleh = pengguna
        change_request.diproses_pada = timezone.now()
        change_request.save(update_fields=['status', 'diproses_oleh', 'diproses_pada'])
        transaction.on_commit(lambda: send_schedule_update(
            change_request.jadwal,
            event='schedule.change_rejected',
        ))
        messages.success(request, 'Permintaan perubahan jadwal ditolak.')
        return redirect('jadwal:jadwal_list')

    jadwal = change_request.jadwal
    jadwal.mata_kuliah = str(change_request.matkul)
    jadwal.kelas = change_request.matkul.kelas
    jadwal.pengampu = change_request.matkul.dosen
    jadwal.ruangan = change_request.ruangan
    jadwal.ruangan_tambahan = change_request.ruangan_tambahan
    jadwal.hari = change_request.hari
    jadwal.waktu_mulai = change_request.waktu_mulai
    jadwal.waktu_selesai = change_request.waktu_selesai
    jadwal.catatan = change_request.catatan
    jadwal.status = JadwalPraktikum.STATUS_DITERIMA
    try:
        jadwal.full_clean()
    except ValidationError as error:
        messages.error(request, f'Perubahan belum dapat disetujui: {error}')
        return redirect('jadwal:jadwal_list')

    jadwal.save()
    change_request.status = 'diterima'
    change_request.diproses_oleh = pengguna
    change_request.diproses_pada = timezone.now()
    change_request.save(update_fields=['status', 'diproses_oleh', 'diproses_pada'])
    transaction.on_commit(lambda: send_schedule_update(jadwal, event='schedule.change_accepted'))
    messages.success(request, 'Perubahan jadwal disetujui dan jadwal telah diperbarui.')
    return redirect('jadwal:jadwal_detail', pk=jadwal.pk)


def available_rooms(request):
    try:
        matkul = MataKuliahAsleb.objects.get(pk=request.GET.get('matkul'))
    except (ValueError, TypeError, MataKuliahAsleb.DoesNotExist):
        return JsonResponse({'participant_count': 0, 'rooms': []})

    participant_count = matkul.peserta_praktikum.filter(aktif=True).count()
    rooms = (
        RuanganLab.objects.filter(aktif=True)
        .filter(Q(kapasitas__isnull=False) | Q(kapasitas_tak_terbatas=True))
        .order_by('kapasitas_tak_terbatas', 'kapasitas', 'nama')
    )
    pengguna = getattr(request, 'current_pengguna', None)
    groups = GrupRuanganGabungan.objects.filter(aktif=True).prefetch_related('ruangan')
    if not participant_count and pengguna and pengguna.role == 'asisten_lab':
        rooms = rooms.none()
    combinable_rooms = {}
    for group in groups:
        grouped_rooms = [room for room in group.ruangan.all() if room.aktif]
        for room in grouped_rooms:
            combinable_rooms[str(room.pk)] = [
                {
                    'id': other_room.pk,
                    'label': str(other_room),
                    'capacity': other_room.kapasitas,
                    'unlimited': other_room.kapasitas_tak_terbatas,
                }
                for other_room in grouped_rooms
                if other_room.pk != room.pk
            ]

    return JsonResponse({
        'participant_count': participant_count,
        'rooms': [
            {
                'id': room.pk,
                'label': str(room),
                'capacity': room.kapasitas,
                'unlimited': room.kapasitas_tak_terbatas,
            }
            for room in rooms
        ],
        'combinable_rooms': combinable_rooms,
    })

